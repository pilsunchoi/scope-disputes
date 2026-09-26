"""두 코더의 조사범위 문언 코딩을 대조해 일치율과 코헨의 카파를 내고 불일치 목록을 만든다.

첫 번째 코더는 Claude 초벌(data/coding/coder-a.csv), 두 번째 코더는 사람이 채운 양식이다.
두 번째 코더의 양식(19_build_coder_form.py가 만든 것)을 다 채운 뒤 data/coding/coder-b.xlsx로 두고 돌린다.
유형은 다섯 유형 × (부과대상 조건, 제외 요건) 열 개의 예·아니오로 대조한다. 요약 유형, 제외 위치,
판 사이 변경, 용도 여부는 범주로 대조하고, 제외 항목 수는 완전 일치와 ±1 일치를 함께 낸다.

    python scripts/20_compare_coders.py [두 번째 코더 파일]

입력:  data/coding/coder-a.csv, data/coding/coder-b.xlsx
산출:  data/coding/agreement.csv       칸별 일치율, 카파, 비교한 조치 수
       data/coding/disagreements.csv   조치 × 칸 불일치 (협의용. resolved 칸은 비워 둔다)
"""
import csv
import sys
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
A = ROOT / "data" / "coding" / "coder-a.csv"
B_DEFAULT = ROOT / "data" / "coding" / "coder-b.xlsx"
OUT_AGREE = ROOT / "data" / "coding" / "agreement.csv"
OUT_DIS = ROOT / "data" / "coding" / "disagreements.csv"

TYPES = ["규격", "가공상태", "용도", "분류", "제품지정"]
CATEGORICAL = ["scope_positive_summary", "exclusion_summary", "exclusion_location", "exclusion_present",
               "changed_across_versions", "use_based"]


def summary(types):
    return "없음" if not types else (types[0] if len(types) == 1 else "복합")


def load_b(path):
    """양식의 코딩 시트. 수식 칸은 엑셀이 저장한 값이 없을 수 있으므로 입력 칸에서 다시 계산한다."""
    ws = load_workbook(path, data_only=True)["코딩"]
    head = [c.value for c in ws[1]]
    rows = {}
    for vals in ws.iter_rows(min_row=2, values_only=True):
        r = dict(zip(head, vals))
        if not r.get("measure_id"):
            continue
        g = lambda k: ("" if r.get(k) is None else str(r[k]).strip())
        pos = [t for t in TYPES if g(f"부과_{t}") == "Y"]
        exc = [t for t in TYPES if g(f"제외_{t}") == "Y"]
        loc = g("exclusion_location")
        rows[g("measure_id")] = {
            **{f"부과_{t}": "Y" if t in pos else "N" for t in TYPES},
            **{f"제외_{t}": "Y" if t in exc else "N" for t in TYPES},
            "scope_positive_summary": summary(pos), "exclusion_summary": summary(exc),
            "exclusion_location": loc, "exclusion_present": "" if not loc else ("N" if loc == "없음" else "Y"),
            "changed_across_versions": g("changed_across_versions") or "N",
            "use_based": "Y" if "용도" in pos + exc else "N",
            "n_exclusion_items": g("n_exclusion_items"),
            "coded": bool(loc or pos or exc),
        }
    return rows


def load_a():
    rows = {}
    for r in csv.DictReader(A.open(encoding="utf-8-sig")):
        pos = [t for t in r["scope_positive_types"].split(";") if t]
        exc = [t for t in r["exclusion_types"].split(";") if t]
        rows[r["measure_id"]] = {
            **{f"부과_{t}": "Y" if t in pos else "N" for t in TYPES},
            **{f"제외_{t}": "Y" if t in exc else "N" for t in TYPES},
            **{k: r[k] for k in CATEGORICAL}, "n_exclusion_items": r["n_exclusion_items"],
        }
    return rows


def kappa(pairs):
    n = len(pairs)
    if n == 0:
        return None, None
    po = sum(a == b for a, b in pairs) / n
    ca, cb = Counter(a for a, _ in pairs), Counter(b for _, b in pairs)
    pe = sum(ca[k] * cb[k] for k in set(ca) | set(cb)) / (n * n)
    return po, (None if pe == 1 else (po - pe) / (1 - pe))


def main():
    path = Path(sys.argv[1]) if len(sys.argv) > 1 else B_DEFAULT
    a, b = load_a(), load_b(path)
    ids = sorted(m for m in a if m in b and b[m]["coded"])
    print(f"두 코더가 모두 코딩한 조치 {len(ids)}개 (두 번째 코더 파일의 행 {len(b)}개)")

    agree, dis = [], []
    cols = [f"부과_{t}" for t in TYPES] + [f"제외_{t}" for t in TYPES] + CATEGORICAL
    for col in cols:
        pairs = [(a[m][col], b[m][col]) for m in ids]
        po, k = kappa(pairs)
        agree.append({"column": col, "n": len(pairs), "agreement": f"{po:.3f}" if po is not None else "",
                      "kappa": f"{k:.3f}" if k is not None else ""})
        dis += [{"measure_id": m, "column": col, "coder_a": a[m][col], "coder_b": b[m][col], "resolved": ""}
                for m in ids if a[m][col] != b[m][col]]
    nums = [(a[m]["n_exclusion_items"], b[m]["n_exclusion_items"]) for m in ids]
    nums = [(int(x), int(y)) for x, y in nums if x.isdigit() and y.isdigit()]
    if nums:
        agree.append({"column": "n_exclusion_items(완전)", "n": len(nums),
                      "agreement": f"{sum(x == y for x, y in nums) / len(nums):.3f}", "kappa": ""})
        agree.append({"column": "n_exclusion_items(±1)", "n": len(nums),
                      "agreement": f"{sum(abs(x - y) <= 1 for x, y in nums) / len(nums):.3f}", "kappa": ""})
    dis += [{"measure_id": m, "column": "n_exclusion_items", "coder_a": a[m]["n_exclusion_items"],
             "coder_b": b[m]["n_exclusion_items"], "resolved": ""}
            for m in ids if a[m]["n_exclusion_items"] != b[m]["n_exclusion_items"]]

    with OUT_AGREE.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=["column", "n", "agreement", "kappa"])
        w.writeheader()
        w.writerows(agree)
    with OUT_DIS.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=["measure_id", "column", "coder_a", "coder_b", "resolved"])
        w.writeheader()
        w.writerows(sorted(dis, key=lambda d: (d["measure_id"], d["column"])))
    for g in agree:
        print(f"  {g['column']:<26} 일치 {g['agreement']:>6}  카파 {g['kappa']:>6}")
    print(f"불일치 {len(dis)}칸 (조치 {len({d['measure_id'] for d in dis})}개) -> {OUT_DIS.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
