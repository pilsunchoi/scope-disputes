"""두 번째 코더(사람)용 조사범위 문언 코딩 양식(엑셀)을 만든다.

코더 간 독립성을 지키려고 양식에는 문언만 넣고 첫 번째 코더(Claude 초벌, coder-a.csv)의 값은 넣지 않는다.
유형은 조치마다 부과대상 조건과 제외 요건 각각 다섯 칸(규격·가공상태·용도·분류·제품지정)에 Y를 고르게 하고,
요약 유형·용도 여부·제외 유무는 수식으로 계산한다. 코더가 채울 칸은 노란색이다.

    python scripts/19_build_coder_form.py

입력:  data/coding/input/batch-*.jsonl (04_build_coding_input.py 산출), data/coding/codebook.md
산출:  data/coding/coder-b-form.xlsx
"""
import glob
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[1]
INPUT = ROOT / "data" / "coding" / "input"
CODEBOOK = ROOT / "data" / "coding" / "codebook.md"
OUT = ROOT / "data" / "coding" / "coder-b-form.xlsx"

FONT = "Malgun Gothic"  # 한글이 들어가는 양식이므로 맑은 고딕을 쓴다
TYPES = ["규격", "가공상태", "용도", "분류", "제품지정"]
F_TEXT = PatternFill("solid", fgColor="F2F2F2")    # 읽기 전용 문언
F_INPUT = PatternFill("solid", fgColor="FFFF99")   # 코더가 채우는 칸
F_CALC = PatternFill("solid", fgColor="DDEBF7")    # 수식으로 계산되는 칸
F_HEAD = PatternFill("solid", fgColor="404040")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

# (머리글, 너비, 종류, 설명)
TEXT_COLS = [
    ("measure_id", 9, "text", "조치 번호"),
    ("제명", 30, "text", "코딩 대상 판의 부령 제명"),
    ("조치 시행일", 11, "text", "조치(첫 판) 시행일"),
    ("코딩 대상 판", 16, "text", "마지막 유효판의 시행일과 제정·개정 구분"),
    ("부과대상 물품 조문", 60, "text", "부령 본문의 부과대상 물품 조문"),
    ("본문 제외 문언", 40, "text", "부과대상 조문 가운데 '제외'를 말하는 부분(자동 추출이라 부과조건이 섞일 수 있음)"),
    ("제외 물품 별표", 60, "text", "제목에 '제외'가 든 별표의 본문"),
    ("부과대상을 정한 별표", 40, "text", "본문이 '별표에 규정된 품목'으로 부과대상을 정할 때의 그 별표(세율표 등)"),
    ("첫 판 부과대상 조문", 40, "text", "첫 판과 문언이 다를 때만. 판 사이 변경(changed) 판단에 쓴다"),
    ("첫 판 제외 별표", 40, "text", "첫 판과 문언이 다를 때만"),
]
IN_TYPES = [(f"부과_{t}", 7, "yn", f"부과대상 조건에 {t} 유형이 있으면 Y") for t in TYPES] + \
           [(f"제외_{t}", 7, "yn", f"제외 요건에 {t} 유형이 있으면 Y") for t in TYPES]
IN_OTHER = [
    ("exclusion_location", 11, "loc", "본문 / 별표 / 둘다 / 없음 (코드북 4절)"),
    ("n_exclusion_items", 9, "int", "제외 항목 수(코드북 3절·5절 5번)"),
    ("changed_across_versions", 11, "yn2", "첫 판과 범위 내용이 다르면 Y, 같으면 N (코드북 5절 4번)"),
    ("evidence_positive", 30, "free", "부과대상 조건의 근거 문구(짧게 인용)"),
    ("evidence_exclusion", 30, "free", "제외 요건의 근거 문구(짧게 인용)"),
    ("confidence", 9, "conf", "높음 / 중간 / 낮음"),
    ("note", 30, "free", "판단이 애매한 이유 등"),
]
CALC = [
    ("scope_positive_types", 18, "부과_ 다섯 칸에서 계산"),
    ("scope_positive_summary", 11, "없음 / 한 유형 / 복합"),
    ("exclusion_types", 18, "제외_ 다섯 칸에서 계산"),
    ("exclusion_summary", 11, "없음 / 한 유형 / 복합"),
    ("exclusion_present", 9, "exclusion_location이 없음이 아니면 Y"),
    ("use_based", 9, "부과_용도나 제외_용도가 Y면 Y"),
]


def records():
    recs = [json.loads(l) for f in sorted(INPUT.glob("batch-*.jsonl")) for l in f.open(encoding="utf-8")]
    return sorted(recs, key=lambda r: r["measure_id"])


def type_list(first_col, row):
    """다섯 칸의 Y를 코드북 순서대로 세미콜론으로 잇는 수식(엑셀 2007 함수만)."""
    parts = [f'IF({get_column_letter(first_col + i)}{row}="Y","{t};","")' for i, t in enumerate(TYPES)]
    joined = "&".join(parts)
    return f'=IF(LEN({joined})=0,"",LEFT({joined},LEN({joined})-1))'


def summary(first_col, row):
    rng = f"{get_column_letter(first_col)}{row}:{get_column_letter(first_col + 4)}{row}"
    one = "".join(f'IF({get_column_letter(first_col + i)}{row}="Y","{t}",' for i, t in enumerate(TYPES)) + '""' + ")" * 5
    return f'=IF(COUNTIF({rng},"Y")=0,"없음",IF(COUNTIF({rng},"Y")>1,"복합",{one}))'


def build_coding(ws, recs):
    cols = TEXT_COLS + IN_TYPES + IN_OTHER + [(h, w, "calc", d) for h, w, d in CALC]
    for j, (h, w, kind, desc) in enumerate(cols, 1):
        c = ws.cell(1, j, h)
        c.font = Font(name=FONT, bold=True, color="FFFFFF", size=9)
        c.fill = F_HEAD
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        c.border = BORDER
        c.comment = Comment(desc, "코딩 양식")
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.row_dimensions[1].height = 42

    pos = {h: j for j, (h, *_rest) in enumerate(cols, 1)}
    p0, e0 = pos["부과_규격"], pos["제외_규격"]
    n = len(recs)
    for i, r in enumerate(recs, 2):
        cv = r["coded_version"]
        vals = [r["measure_id"], r["decree_title"], r["effective_date"],
                f'{cv["effective_date"]} {cv["revision_type"]}', r["scope_text"], r["body_exclusion_text"],
                r["annex_exclusion_text"], r.get("annex_other_text", ""),
                r.get("first_version_scope_text", ""), r.get("first_version_annex_exclusion_text", "")]
        for j, v in enumerate(vals, 1):
            c = ws.cell(i, j, v)
            c.fill = F_TEXT
        for h, *_rest in IN_TYPES + IN_OTHER:
            ws.cell(i, pos[h]).fill = F_INPUT
        ws.cell(i, pos["scope_positive_types"], type_list(p0, i))
        ws.cell(i, pos["scope_positive_summary"], summary(p0, i))
        ws.cell(i, pos["exclusion_types"], type_list(e0, i))
        ws.cell(i, pos["exclusion_summary"], summary(e0, i))
        loc = f'{get_column_letter(pos["exclusion_location"])}{i}'
        ws.cell(i, pos["exclusion_present"], f'=IF({loc}="","",IF({loc}="없음","N","Y"))')
        pu, eu = f"{get_column_letter(p0 + 2)}{i}", f"{get_column_letter(e0 + 2)}{i}"
        ws.cell(i, pos["use_based"], f'=IF(OR({pu}="Y",{eu}="Y"),"Y","N")')
        for h, *_rest in CALC:
            ws.cell(i, pos[h]).fill = F_CALC
        for j in range(1, len(cols) + 1):
            c = ws.cell(i, j)
            c.font = Font(name=FONT, size=9)
            c.alignment = Alignment(wrap_text=True, vertical="top")
            c.border = BORDER
        ws.row_dimensions[i].height = 110

    last = n + 1
    def dv(formula, cols_, msg):
        v = DataValidation(type="list", formula1=formula, allow_blank=True, showErrorMessage=True,
                           errorTitle="값 확인", error=msg)
        ws.add_data_validation(v)
        for h in cols_:
            L = get_column_letter(pos[h])
            v.add(f"{L}2:{L}{last}")
    dv('"Y"', [h for h, *_ in IN_TYPES], "Y만 넣거나 비워 두세요")
    dv('"본문,별표,둘다,없음"', ["exclusion_location"], "본문·별표·둘다·없음 가운데 하나")
    dv('"Y,N"', ["changed_across_versions"], "Y 또는 N")
    dv('"높음,중간,낮음"', ["confidence"], "높음·중간·낮음 가운데 하나")
    iv = DataValidation(type="whole", operator="between", formula1="0", formula2="500", allow_blank=True,
                        showErrorMessage=True, errorTitle="값 확인", error="0 이상의 정수")
    ws.add_data_validation(iv)
    L = get_column_letter(pos["n_exclusion_items"])
    iv.add(f"{L}2:{L}{last}")

    # 제외 위치가 '없음'이 아닌데 제외 유형이 하나도 없으면 빨갛게 표시한다
    loc_L = get_column_letter(pos["exclusion_location"])
    er = f"{get_column_letter(e0)}2:{get_column_letter(e0 + 4)}2"
    ws.conditional_formatting.add(
        f"{loc_L}2:{loc_L}{last}",
        FormulaRule(formula=[f'AND({loc_L}2<>"",{loc_L}2<>"없음",COUNTIF({er},"Y")=0)'],
                    fill=PatternFill("solid", fgColor="F4B6B6")))
    ws.freeze_panes = ws.cell(2, 2)
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{last}"
    return pos


def build_guide(ws, n):
    rows = [
        ("조사범위 문언 코딩 양식 — 두 번째 코더용", True),
        ("", False),
        ("이 양식은 무역구제 조사범위 분쟁 DB의 1단계 코딩을 두 번째 코더가 독립으로 하기 위한 것이다.", False),
        (f"'코딩' 시트에 조치 {n}개가 한 행씩 있다. 코딩 기준은 '코드북' 시트(data/coding/codebook.md와 같은 내용)이며, 5절의 결정까지 포함해 따른다.", False),
        ("독립성: 첫 번째 코더의 값(data/coding/coder-a.csv)은 코딩이 끝날 때까지 열어 보지 않는다.", False),
        ("", False),
        ("칸의 색", True),
        ("회색  문언(읽기 전용). 부령 본문과 별표에서 자동으로 뽑은 것이다. 본문 제외 문언에는 부과조건이 섞일 수 있으니 부과대상 조문과 함께 읽는다.", False),
        ("노란색  코더가 채우는 칸. 목록에서 고르거나(Y, 본문/별표/둘다/없음, Y/N, 높음/중간/낮음) 숫자·문구를 적는다.", False),
        ("파란색  수식으로 계산되는 칸. 고치지 않는다. 유형 목록, 요약 유형(없음/한 유형/복합), 제외 유무, 용도 여부가 여기서 나온다.", False),
        ("빨간색  제외 위치를 '없음'이 아닌 값으로 두었는데 제외 유형이 하나도 없을 때 표시된다.", False),
        ("", False),
        ("채우는 순서", True),
        ("1. 부과대상 물품 조문과 부과대상을 정한 별표를 읽고, 물품명과 기본 HSK 열거 말고 부과대상을 좁히는 조건의 유형에 Y를 넣는다(부과_ 다섯 칸).", False),
        ("2. 본문 제외 문언과 제외 물품 별표를 읽고, 제외 요건의 유형에 Y를 넣는다(제외_ 다섯 칸). 공급자를 빼는 별표는 제외 요건이 아니다.", False),
        ("3. exclusion_location, n_exclusion_items를 채운다.", False),
        ("4. 첫 판 문언 칸이 채워져 있으면 범위의 내용이 바뀌었는지 보고 changed_across_versions를 고른다. 비어 있으면 N이다.", False),
        ("5. 근거 문구, 확신도, 비고를 적는다.", False),
        ("", False),
        ("다 채운 뒤 파일 이름을 coder-b.xlsx로 바꾸어 data/coding/에 두면 scripts/20_compare_coders.py가 두 코더를 대조한다.", False),
        ("", False),
        ("예시 한 행 (값의 형식만 보여 주려는 가상의 예이며 '코딩' 시트의 어느 조치와도 무관하다)", True),
    ]
    for i, (t, bold) in enumerate(rows, 1):
        c = ws.cell(i, 1, t)
        c.font = Font(name=FONT, bold=bold, size=12 if i == 1 else 10)
        c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 120
    ex_head = ["부과_규격", "부과_가공상태", "제외_가공상태", "제외_용도", "exclusion_location", "n_exclusion_items",
               "changed_across_versions", "evidence_positive", "evidence_exclusion", "confidence", "note"]
    ex_val = ["Y", "Y", "Y", "Y", "별표", 4, "N", "두께가 6밀리미터 이상인 것", "별표 1: 모터 절연용 흰색 필름 등",
              "중간", "별표 3호가 용도와 색을 함께 적어 두 유형으로 봄"]
    r0 = len(rows) + 1
    for j, (h, v) in enumerate(zip(ex_head, ex_val)):
        hc = ws.cell(r0, 2 + j, h)
        hc.font = Font(name=FONT, bold=True, size=9)
        vc = ws.cell(r0 + 1, 2 + j, v)
        vc.font = Font(name=FONT, size=9)
        vc.fill = F_INPUT
        ws.column_dimensions[get_column_letter(2 + j)].width = 16


def build_codebook(ws):
    ws.column_dimensions["A"].width = 150
    for i, line in enumerate(CODEBOOK.read_text(encoding="utf-8").splitlines(), 1):
        c = ws.cell(i, 1, line)
        c.font = Font(name=FONT, size=10, bold=line.startswith("#"))
        c.alignment = Alignment(wrap_text=True, vertical="top")


def main():
    recs = records()
    wb = Workbook()
    guide = wb.active
    guide.title = "안내"
    build_guide(guide, len(recs))
    build_coding(wb.create_sheet("코딩"), recs)
    build_codebook(wb.create_sheet("코드북"))
    wb.save(OUT)
    print(f"조치 {len(recs)}개 -> {OUT.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
